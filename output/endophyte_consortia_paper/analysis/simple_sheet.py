import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

F = 'Times New Roman'
YEL = PatternFill('solid', fgColor='FFF2CC'); HDR = PatternFill('solid', fgColor='D9E2F3')
thin = Side(style='thin', color='999999'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def ft(b=False, i=False): return Font(name=F, bold=b, italic=i, size=11)

INOC = ['Control', 'IS-06', 'IS-05', 'C4-5']; SAL = [0, 50, 100, 150]; REPS = 4
rows = [(i, s, r) for i in INOC for s in SAL for r in range(1, REPS + 1)]
random.seed(2026); pos = list(range(1, len(rows) + 1)); random.shuffle(pos)

# (header, kind) kind: 'in' input, or formula template using column letters of named inputs
COLS = [
 ('Pot ID', 'id'), ('Treatment', 'id'), ('NaCl (mM)', 'id'), ('Replicate', 'id'), ('Bench position', 'id'),
 ('Plant height (cm)', 'in'), ('No. of leaves', 'in'), ('Leaf area (cm²/plant)', 'in'),
 ('Shoot fresh weight (g)', 'in'), ('Shoot dry weight (g)', 'in'),
 ('Root length (cm)', 'in'), ('Root fresh weight (g)', 'in'), ('Root dry weight (g)', 'in'),
 ('Root/shoot ratio', ('Root dry weight (g)', 'Shoot dry weight (g)', '{0}/{1}')),
 ('Root volume (cm³)', 'in'), ('Root surface area (cm²)', 'in'),
 ('Relative water content (%)', 'in'), ('Electrolyte leakage (%)', 'in'),
 ('Chlorophyll a (mg/g FW)', 'in'), ('Chlorophyll b (mg/g FW)', 'in'),
 ('Total chlorophyll (mg/g FW)', ('Chlorophyll a (mg/g FW)', 'Chlorophyll b (mg/g FW)', '{0}+{1}')),
 ('Carotenoids (mg/g FW)', 'in'),
 ('Proline (µmol/g FW)', 'in'), ('Soluble sugars (mg/g FW)', 'in'), ('Soluble protein (mg/g FW)', 'in'),
 ('MDA (nmol/g FW)', 'in'), ('H₂O₂ (µmol/g FW)', 'in'),
 ('SOD (U/mg protein)', 'in'), ('CAT (µmol/min/mg protein)', 'in'), ('POD (µmol/min/mg protein)', 'in'), ('APX (µmol/min/mg protein)', 'in'),
 ('Leaf Na⁺ (mg/g DW)', 'in'), ('Leaf K⁺ (mg/g DW)', 'in'),
 ('Leaf K⁺/Na⁺ ratio', ('Leaf K⁺ (mg/g DW)', 'Leaf Na⁺ (mg/g DW)', '({0}/39.1)/({1}/22.99)')),
 ('Root Na⁺ (mg/g DW)', 'in'), ('Root K⁺ (mg/g DW)', 'in'),
 ('Root K⁺/Na⁺ ratio', ('Root K⁺ (mg/g DW)', 'Root Na⁺ (mg/g DW)', '({0}/39.1)/({1}/22.99)')),
 ('Leaf Ca²⁺ (mg/g DW)', 'in'), ('Leaf Mg²⁺ (mg/g DW)', 'in'),
 ('Remarks', 'in'),
]
idx = {h: i + 1 for i, (h, _) in enumerate(COLS)}

wb = Workbook()
# ---------- Sheet 1: Data ----------
ws = wb.active; ws.title = 'Data'
ws['A1'] = 'Rice cv. IR-64 | Endophyte inoculation × NaCl | 4 treatments × 4 NaCl levels × 4 replicates = 64 pots | Harvest at 45 DAS'
ws['A1'].font = ft(True)
ws['A2'] = 'Fill the shaded cells (one row per pot; mean of the 3 plants in the pot). Leave a cell blank if not measured. Grey columns calculate automatically.'
ws['A2'].font = ft(i=True)
H = 4
for h, k in COLS:
    c = ws.cell(H, idx[h], h); c.font = ft(True); c.fill = HDR; c.border = BOX
    c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
ws.row_dimensions[H].height = 60
for n, (inoc, s, r) in enumerate(rows):
    R = H + 1 + n
    vals = {'Pot ID': f'P{n+1:02d}', 'Treatment': inoc, 'NaCl (mM)': s, 'Replicate': r, 'Bench position': pos[n]}
    for h, k in COLS:
        c = ws.cell(R, idx[h]); c.border = BOX; c.font = ft()
        if k == 'id': c.value = vals[h]; c.alignment = Alignment(horizontal='center')
        elif k == 'in': c.fill = YEL
        else:
            a, b, f = k; ra, rb = f'{L(idx[a])}{R}', f'{L(idx[b])}{R}'
            c.value = f'=IF(OR({ra}="",{rb}=""),"",{f.format(ra, rb)})'
            c.number_format = '0.00'; c.fill = PatternFill('solid', fgColor='EDEDED')
for h, i in idx.items(): ws.column_dimensions[L(i)].width = 9 if i <= 5 else 13
ws.column_dimensions['B'].width = 10; ws.column_dimensions[L(idx['Remarks'])].width = 30
ws.freeze_panes = ws.cell(H + 1, 3)
last = H + len(rows)

# ---------- Sheet 2: Summary ----------
sm = wb.create_sheet('Summary')
sm['A1'] = 'Treatment means (calculated automatically from the Data sheet)'; sm['A1'].font = ft(True)
traits = [h for h, k in COLS if k != 'id' and h != 'Remarks']
sm.cell(3, 1, 'Treatment').font = ft(True); sm.cell(3, 2, 'NaCl (mM)').font = ft(True)
for j, t in enumerate(traits, 3):
    c = sm.cell(3, j, t); c.font = ft(True); c.fill = HDR; c.alignment = Alignment(wrap_text=True, horizontal='center'); c.border = BOX
    sm.column_dimensions[L(j)].width = 12
sm.cell(3, 1).fill = HDR; sm.cell(3, 2).fill = HDR; sm.row_dimensions[3].height = 60
r = 4
for inoc in INOC:
    for s in SAL:
        sm.cell(r, 1, inoc).font = ft(); sm.cell(r, 2, s).font = ft()
        for j, t in enumerate(traits, 3):
            col = L(idx[t])
            c = sm.cell(r, j, f'=IFERROR(AVERAGEIFS(Data!${col}$5:${col}${last},Data!$B$5:$B${last},$A{r},Data!$C$5:$C${last},$B{r}),"")')
            c.number_format = '0.00'; c.font = ft(); c.border = BOX
        r += 1
sm.freeze_panes = 'C4'

# ---------- Sheet 3: Methods ----------
mt = wb.create_sheet('Methods')
mt['A1'] = 'Short method and unit guide (calculate each value with the lab formula, then enter the final value in the Data sheet)'; mt['A1'].font = ft(True)
guide = [
 ('NaCl solutions', '50, 100, 150 mM = 2.922, 5.844, 8.766 g NaCl per litre; 500 mL per pot on alternate days for 30 days. Note the exact start day.'),
 ('Dry weight', 'Oven 70 °C for 72 h, to constant weight.'),
 ('Relative water content', '(Fresh wt − Dry wt) / (Turgid wt − Dry wt) × 100; turgid wt after 4 h floating on water in the dark.'),
 ('Electrolyte leakage', 'EC1 (leaf discs in water, 2 h) / EC2 (after boiling 20 min) × 100.'),
 ('Chlorophyll / carotenoids', 'DMSO extraction at 60 °C; Chl a = 12.7 A663 − 2.69 A645; Chl b = 22.9 A645 − 4.68 A663 (mg/L) × V / (1000 × W). Carotenoids at 480 nm.'),
 ('Proline', 'Acid-ninhydrin, toluene layer at 520 nm, proline standard curve.'),
 ('Soluble sugars / protein', 'Anthrone at 620 nm (glucose standard); Lowry at 660 nm (BSA standard).'),
 ('MDA', '[(A532 − A600) / 155] × 1000 × extract volume (mL) / fresh weight (g) = nmol/g FW.'),
 ('H₂O₂', 'KI method at 390 nm, H₂O₂ standard curve.'),
 ('SOD', 'NBT method at 560 nm; 1 unit = 50% inhibition; divide by protein (mg).'),
 ('CAT / APX / POD', 'Rate per minute at 240 / 290 / 470 nm; ε = 0.0436 / 2.8 / 26.6 mM⁻¹ cm⁻¹; divide by protein (mg).'),
 ('Na⁺, K⁺, Ca²⁺, Mg²⁺', 'Dry tissue digested or ashed; Na⁺ and K⁺ by flame photometer; Ca²⁺ and Mg²⁺ by EDTA titration or AAS. mg/g DW = reading (mg/L) × volume (mL) / 1000 × dilution / dry weight (g).'),
 ('Important', 'Enter every pot separately (not only averages). Keep the raw readings in your notebook.'),
]
for i, (a, b) in enumerate(guide, 3):
    mt.cell(i, 1, a).font = ft(True); c = mt.cell(i, 2, b); c.font = ft(); c.alignment = Alignment(wrap_text=True, vertical='top')
mt.column_dimensions['A'].width = 26; mt.column_dimensions['B'].width = 110

wb.save('Data_Sheet_Simple.xlsx'); print('saved', len(rows), 'pots,', len(traits), 'traits')
