import re, json, sys, pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from refs import R

BLACK = RGBColor(0, 0, 0)
FONT = 'Times New Roman'
OUT = sys.argv[1] if len(sys.argv) > 1 else 'manuscript.docx'

text = open('ms_part1.txt').read() + '\n\n' + open('ms_part2.txt').read() + '\n\n' + open('ms_part3.txt').read()

# ---------------- citation audit ----------------
body_for_cites = re.sub(r'\[\[.*?\]\]', '', text)
def ref_patterns(r):
    cite = r[8]
    m = re.match(r'(.+), (\d{4})$', cite)
    name, yr = m.group(1), m.group(2)
    pats = [re.escape(name) + r',\s*' + yr]                       # (Name, 2020) / (Name et al., 2020) / (A & B, 2020)
    narr = name.replace(' & ', ' and ')
    pats.append(re.escape(narr) + r"(?:'s)?\s*\(" + yr + r'\)')     # Name and Name (2020) / Name et al. (2020) / Arnon's (1949)
    return pats
missing = []
for r in R:
    if not any(re.search(p, body_for_cites) for p in ref_patterns(r)):
        missing.append(r[8])
# every citation-looking token must match a reference
cite_tokens = set()
for grp in re.findall(r'\(([^()]*?\d{4}[^()]*?)\)', body_for_cites):
    for part in grp.split(';'):
        part = part.strip()
        if re.search(r'[A-Z][^,]*,\s*\d{4}$', part):
            cite_tokens.add(part)
keys = {r[8] for r in R}
orphans = sorted(t for t in cite_tokens if t not in keys)
narr = set(re.findall(r"([A-Z][A-Za-zÀ-ž\-]+(?: et al\.| and [A-Z][A-Za-zÀ-ž\-]+)?(?:'s)?) \((\d{4})\)", body_for_cites))
narr_bad = []
for n, y in narr:
    n2 = n.replace("'s", '').replace(' and ', ' & ')
    if not any(k.endswith(f'{n2}, {y}') for k in keys):
        narr_bad.append(f'{n} ({y})')
print('Uncited references:', missing)
print('Parenthetical citations without reference:', orphans)
print('Narrative citations without reference:', narr_bad)
if missing or orphans or narr_bad:
    sys.exit('CITATION AUDIT FAILED')

# ---------------- helpers ----------------
def set_run(run, italic=None, bold=None, sup=False, sub=False, size=12):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = BLACK
    if italic is not None: run.italic = italic
    if bold is not None: run.bold = bold
    if sup: run.font.superscript = True
    if sub: run.font.subscript = True

TOK = re.compile(r'(\*\*.+?\*\*|\*.+?\*|<i>.+?</i>|\^\{.+?\}|_\{.+?\})')
def add_rich(par, s, bold=False, italic=False, size=12):
    for piece in TOK.split(s):
        if not piece: continue
        if piece.startswith('**'):
            add_rich(par, piece[2:-2], bold=True, italic=italic, size=size)
        elif piece.startswith('<i>'):
            add_rich(par, piece[3:-4], bold=bold, italic=not italic, size=size)
        elif piece.startswith('*'):
            add_rich(par, piece[1:-1], bold=bold, italic=not italic, size=size)
        elif piece.startswith('^{'):
            r = par.add_run(piece[2:-1]); set_run(r, italic=italic, bold=bold, sup=True, size=size)
        elif piece.startswith('_{'):
            r = par.add_run(piece[2:-1]); set_run(r, italic=italic, bold=bold, sub=True, size=size)
        else:
            r = par.add_run(piece); set_run(r, italic=italic, bold=bold, size=size)

def fmt(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, spacing=1.5, first=None, hanging=None, keep=False):
    pf = par.paragraph_format
    par.alignment = align
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.line_spacing = spacing
    if first is not None: pf.first_line_indent = Cm(first)
    if hanging is not None:
        pf.left_indent = Cm(hanging); pf.first_line_indent = Cm(-hanging)
    if keep: pf.keep_with_next = True

def para(doc, s, **kw):
    p = doc.add_paragraph(); add_rich(p, s); fmt(p, **kw); return p

def heading(doc, s, level):
    p = doc.add_paragraph()
    add_rich(p, s, bold=True, italic=(level == 3))
    fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=12 if level == 1 else 8, after=6, keep=True)
    return p

def hyperlink(par, url, label):
    rid = par.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    f = OxmlElement('w:rFonts'); f.set(qn('w:ascii'), FONT); f.set(qn('w:hAnsi'), FONT); rPr.append(f)
    c = OxmlElement('w:color'); c.set(qn('w:val'), '000000'); rPr.append(c)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24'); rPr.append(sz)
    r.append(rPr); t = OxmlElement('w:t'); t.text = label; t.set(qn('xml:space'), 'preserve'); r.append(t)
    h.append(r); par._p.append(h)

def cell_text(cell, s, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, size=12):
    cell.text = ''
    p = cell.paragraphs[0]; add_rich(p, s, bold=bold, size=size)
    p.alignment = align; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def borders(table, rows_with_bottom):
    # academic three-line style: top rule, rule under header, bottom rule
    tbl = table._tbl
    tblPr = tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge, val in [('top', 'single'), ('bottom', 'single'), ('left', 'nil'), ('right', 'nil'), ('insideH', 'nil'), ('insideV', 'nil')]:
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), val); e.set(qn('w:sz'), '8'); e.set(qn('w:color'), '000000'); b.append(e)
    tblPr.append(b)
    for ri in rows_with_bottom:
        for c in table.rows[ri].cells:
            tcPr = c._tc.get_or_add_tcPr(); tb = OxmlElement('w:tcBorders')
            e = OxmlElement('w:bottom'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '6'); e.set(qn('w:color'), '000000')
            tb.append(e); tcPr.append(tb)

def new_section(doc, landscape):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    else:
        sec.orientation = WD_ORIENT.PORTRAIT; sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(sec, m, Cm(2.54))
    return sec


def fix_widths(t, ws):
    tblPr = t._tbl.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    tw = OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(sum(ws) * 567))); tw.set(qn('w:type'), 'dxa'); tblPr.append(tw)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), ws): gc.set(qn('w:w'), str(int(w * 567)))
    for row in t.rows:
        for c, w in zip(row.cells, ws): c.width = Cm(w)

# ---------------- data for tables ----------------
lab = pd.read_csv('lab_table.csv')
D = json.load(open('derived.json'))
S = json.load(open('stats.json'))
df = pd.read_csv('data_long.csv')

def f1(x): return f'{x:.1f}'
def fF(F):
    if F >= 10000:
        e = int(f'{F:e}'.split('e')[1]); m = F / 10**e
        return f'{m:.2f} × 10^{{{e}}}'
    return f'{F:.1f}'
AST='\u2217'
def star(p): return AST*3 if p < 0.001 else (AST*2 if p < 0.01 else (AST if p < 0.05 else ' ns'))

def table1(doc):
    cap = doc.add_paragraph(); add_rich(cap, '**Table 1.** Germination and early seedling growth of rice cv. IR-64 bacterized with single endophytic isolates or dual-strain consortia in the laboratory assay (0 mM NaCl; 25 seeds per treatment).')
    fmt(cap, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=6, after=4, spacing=1.0, keep=True)
    hdr = ['Treatment', 'Germinated / total', 'Germination (%)', 'Root length (mm)', 'Shoot length (mm)', 'SVI']
    rows = []
    for exp, lab_ in [('lab_iso', 'Single isolates (batch 1)'), ('lab_con', 'Dual-strain consortia (batch 2)')]:
        rows.append(('SEC', lab_))
        for _, r in lab[lab['Unnamed: 0'] == exp].iterrows():
            name = 'Control' if r.inoc == 'Control' else r.inoc
            rows.append((name, f'{int(r.germ_n)} / {int(r.seeds)}', f'{r.G:.0f}', f'{r.root_mm:.2f} ± {r.root_sd:.2f}', f'{r.shoot_mm:.2f} ± {r.shoot_sd:.2f}', f'{r.SVI:.1f}'))
    t = doc.add_table(rows=1 + len(rows), cols=6); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(hdr): cell_text(t.rows[0].cells[i], h, bold=True)
    for ri, row in enumerate(rows, start=1):
        if row[0] == 'SEC':
            c = t.rows[ri].cells[0].merge(t.rows[ri].cells[5]); cell_text(c, f'*{row[1]}*', align=WD_ALIGN_PARAGRAPH.LEFT)
        else:
            for ci, v in enumerate(row): cell_text(t.rows[ri].cells[ci], v, align=WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    for row in t.rows[:-1]:
        for c in row.cells:
            for pp in c.paragraphs: pp.paragraph_format.keep_with_next = True
    borders(t, [0])
    fix_widths(t, [2.5, 2.6, 3.1, 2.9, 2.9, 1.6])
    li = S['lab_iso_chi2']; lc = S['lab_con_chi2']
    note = doc.add_paragraph()
    add_rich(note, f"Root and shoot lengths are means ± SD. SVI = germination (%) × (root + shoot length, cm), recalculated from the recorded values (Abdul-Baki & Anderson, 1973). Germination did not differ among treatments within either batch (isolates: χ^{{2}} = {li['chi2']:.2f}, df = {li['df']}, *p* = {li['p']:.3f}; consortia: χ^{{2}} = {lc['chi2']:.2f}, df = {lc['df']}, *p* = {lc['p']:.3f}).")
    fmt(note, spacing=1.0, before=4, after=10)

def table2(doc):
    cap = doc.add_paragraph(); add_rich(cap, '**Table 2.** Germination (%) and seedling vigour index (SVI) of rice cv. IR-64 bacterized with endophytic isolates or consortia and grown under graded NaCl in the polyhouse (30 seeds per treatment).')
    fmt(cap, before=0, after=4, spacing=1.0, keep=True)
    t = doc.add_table(rows=3 + 12 + 1, cols=9); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    a = t.rows[0].cells[1].merge(t.rows[0].cells[4]); cell_text(a, 'Germination (%)', bold=True)
    b = t.rows[0].cells[5].merge(t.rows[0].cells[8]); cell_text(b, 'Seedling vigour index', bold=True)
    c = t.rows[0].cells[0].merge(t.rows[1].cells[0]); cell_text(c, 'Treatment', bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    for i, s in enumerate(['0 mM', '50 mM', '100 mM', '150 mM'] * 2): cell_text(t.rows[1].cells[i + 1], s, bold=True)
    ri = 2
    def put(name, exp, inoc):
        nonlocal ri
        cells = t.rows[ri].cells
        cell_text(cells[0], name, align=WD_ALIGN_PARAGRAPH.LEFT)
        for j, s in enumerate(['0', '50', '100', '150']):
            cell_text(cells[1 + j], f1(D[exp]['germ'][s][inoc])); cell_text(cells[5 + j], f"{D[exp]['svi'][s][inoc]:.0f}")
        ri += 1
    put('Uninoculated control', 'iso', 'Control')
    sec = t.rows[ri].cells[0].merge(t.rows[ri].cells[8]); cell_text(sec, '*Single isolates*', align=WD_ALIGN_PARAGRAPH.LEFT); ri += 1
    for i in ['IS-01', 'IS-02', 'IS-03', 'IS-04', 'IS-05', 'IS-06']: put(i, 'iso', i)
    sec = t.rows[ri].cells[0].merge(t.rows[ri].cells[8]); cell_text(sec, '*Dual-strain consortia*', align=WD_ALIGN_PARAGRAPH.LEFT); ri += 1
    for i in ['C1-3', 'C1-5', 'C2-3', 'C2-5', 'C4-5']: put(i, 'con', i)
    borders(t, [1])
    fix_widths(t, [5.0] + [2.35] * 8)
    gi, gc = S['iso_germ_glm'], S['con_germ_glm']
    note = doc.add_paragraph()
    add_rich(note, "Germination = germinated seeds / 30 × 100. SVI = germination (%) × (root + shoot length, cm), recalculated from the recorded root and shoot lengths. "
             f"Binomial GLM likelihood-ratio tests, isolate experiment: NaCl χ^{{2}} = {gi['sal']['LR']:.2f} (df 3, *p* < 0.001), inoculant χ^{{2}} = {gi['inoc']['LR']:.2f} (df 6, *p* = {gi['inoc']['p']:.3f}), interaction χ^{{2}} = {gi['inter']['LR']:.2f} (df 18, *p* = {gi['inter']['p']:.3f}); "
             f"consortium experiment: NaCl χ^{{2}} = {gc['sal']['LR']:.2f} (df 3, *p* < 0.001), inoculant χ^{{2}} = {gc['inoc']['LR']:.2f} (df 5, *p* = {gc['inoc']['p']:.3f}), interaction χ^{{2}} = {gc['inter']['LR']:.2f} (df 15, *p* = {gc['inter']['p']:.3f}). The uninoculated control was shared by both experiments.")
    fmt(note, spacing=1.0, before=4, after=6)

def table3(doc):
    cap = doc.add_paragraph(); add_rich(cap, '**Table 3.** Two-way ANOVA (inoculant × NaCl) for seedling growth and leaf biochemical traits of rice cv. IR-64, computed from replicate means and standard deviations (n = 3).')
    fmt(cap, before=0, after=4, spacing=1.0, keep=True)
    rows = [('Root length (cm)', 'root_cm'), ('Shoot length (cm)', 'shoot_cm'), ('Chlorophyll *a* (mg g^{−1} FW)', 'chla'), ('Chlorophyll *b* (mg g^{−1} FW)', 'chlb'),
            ('Soluble protein (µg mL^{−1})', 'protein'), ('Soluble sugars (µg mL^{−1})', 'sugar'), ('Total phenolics (µg mL^{−1})', 'phenol'), ('DPPH scavenging (%)', 'dpph')]
    body = []
    for exp, lab_, dfs in [('iso', 'Single-isolate experiment (7 inoculation levels × 4 NaCl levels)', ('6, 56', '3, 56', '18, 56')),
                           ('con', 'Consortium experiment (6 inoculation levels × 4 NaCl levels)', ('5, 48', '3, 48', '15, 48'))]:
        body.append(('SEC', lab_))
        for name, key in rows:
            k = f'{exp}_{key}_anova'
            if k not in S: continue
            a = S[k]
            body.append((name, fF(a['inoc']['F']) + star(a['inoc']['p']), fF(a['sal']['F']) + star(a['sal']['p']), fF(a['inter']['F']) + star(a['inter']['p']),
                         f"{a['inter']['pe2']:.3f}", f"{a['HSD']:.2f}"))
    t = doc.add_table(rows=1 + len(body), cols=6); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = ['Trait', 'Inoculant F', 'NaCl F', 'Inoculant × NaCl F', 'η_{p}^{2} (interaction)', 'Tukey HSD_{0.05}']
    for i, h in enumerate(hdr): cell_text(t.rows[0].cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    for ri, row in enumerate(body, start=1):
        if row[0] == 'SEC':
            c = t.rows[ri].cells[0].merge(t.rows[ri].cells[5]); cell_text(c, f'*{row[1]}*', align=WD_ALIGN_PARAGRAPH.LEFT)
        else:
            for ci, v in enumerate(row): cell_text(t.rows[ri].cells[ci], v, align=WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    borders(t, [0])
    fix_widths(t, [7.4, 3.4, 3.4, 4.0, 3.5, 2.9])
    note = doc.add_paragraph()
    add_rich(note, "\u2217\u2217\u2217 *p* < 0.001; ns, not significant (*p* ≥ 0.05). η_{p}^{2}, partial eta-squared. HSD is the minimum significant difference between any two treatment means of the same experiment, in the units of the trait. "
             "Degrees of freedom: isolates F_{6,56} (inoculant), F_{3,56} (NaCl), F_{18,56} (interaction); consortia F_{5,48}, F_{3,48}, F_{15,48}. Isolate phenolics, consortium sugars and consortium DPPH are omitted because replicate SDs were not recorded. "
             "Several traits carried identical replicate SDs (±0.15) across treatments, which inflates F; effect sizes and treatment rankings are the more reliable guide.")
    fmt(note, spacing=1.0, before=4, after=6)

FIGS = {
 1: ('Fig1_workflow.png', 14.2, "**Fig. 1.** Infographic flowchart of the experimental workflow. Stage 1, preparation of single-isolate inocula and dual-strain consortia with purity checks; Stage 2, surface sterilization and bacterization of rice cv. IR-64 seeds; Stage 3A, laboratory germination assay; Stage 3B, polyhouse pot trial under 0–150 mM NaCl; Stage 4, leaf biochemical assays on polyhouse plants; Stage 5, statistical analysis, derived indices and multivariate analysis. DAS, days after sowing; CRD, completely randomized design; CMI, composite membership index."),
 2: ('Fig2_germination.png', 15.9, "**Fig. 2.** Germination (%) of rice cv. IR-64 under 0, 50, 100 and 150 mM NaCl in the polyhouse after bacterization with (a) single endophytic isolates and (b) dual-strain consortia. Each bar is the percentage of 30 sown seeds. Salinity reduced germination (*p* < 0.001 in both experiments), whereas inoculant effects were not significant (*p* = 0.071 and 0.065; binomial GLM likelihood-ratio tests)."),
 3: ('Fig3_doseresponse.png', 15.9, "**Fig. 3.** Dose–response of (a) chlorophyll *a*, (b) soluble protein, (c) soluble sugars and (d) root length to NaCl in uninoculated rice and in plants inoculated with the two best single isolates (IS-06, IS-05; solid lines) and the two best consortia (C4-5, C2-5; dashed lines). Error bars are SD (n = 3); replicate SDs were not recorded for consortium sugars. The low IS-05 root value at 100 mM is shown as recorded."),
 4: ('Fig4_heatmap.png', 15.9, "**Fig. 4.** Heat map of inoculation-mediated mitigation (IMM, %), i.e. the change in each trait relative to the uninoculated control at the same NaCl level, for six single isolates (upper block) and five consortia (lower block) at 50, 100 and 150 mM NaCl. Blue cells denote improvement and red cells impairment; the colour scale is truncated at −100% and +300%, and cell labels give exact values. Germ., germination; Car., carotenoids."),
 5: ('Fig5_radar.png', 13.0, "**Fig. 5.** Radar chart of nine traits rescaled to 0–1 across all 48 treatment combinations, comparing unstressed and 150 mM NaCl uninoculated plants with plants inoculated with IS-06 and C4-5 under 150 mM NaCl. A larger polygon indicates a trait profile closer to the best observed values."),
 6: ('Fig6_pca.png', 15.9, "**Fig. 6.** Principal component analysis biplot of nine standardized traits across 48 treatment combinations. Symbol shape and colour indicate inoculation type and symbol size indicates NaCl level; selected treatments are labelled as inoculant–NaCl (mM). Arrows show trait loadings (scaled ×3.2). PC1 represents general seedling growth and metabolite status; PC2 separates protein and sugar content from root and shoot length and DPPH activity."),
}
def figure(doc, n):
    fn, w, cap = FIGS[n]
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(fn, width=Cm(w))
    c = doc.add_paragraph(); add_rich(c, cap); fmt(c, spacing=1.0, before=4, after=12)

# ---------------- document ----------------
doc = Document()
st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(12); st.font.color.rgb = BLACK
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
sec = doc.sections[0]; sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(sec, m, Cm(2.54))

# line numbers (continuous) and page numbers
def line_numbers(section):
    ln = OxmlElement('w:lnNumType'); ln.set(qn('w:countBy'), '1'); ln.set(qn('w:restart'), 'continuous'); ln.set(qn('w:distance'), '283')
    section._sectPr.append(ln)
def page_numbers(section):
    fp = section.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(); set_run(r)
    for kind, txt in [('begin', None), (None, 'PAGE'), ('end', None)]:
        if kind:
            f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), kind); r._r.append(f)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; r._r.append(it)
line_numbers(sec); page_numbers(sec)

blocks = [b.strip('\n') for b in re.split(r'\n\s*\n', text) if b.strip()]
pending_tables = []
for b in blocks:
    if b.startswith('@TITLE'):
        p = doc.add_paragraph(); add_rich(p, b[7:].strip(), bold=True); fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=12, spacing=1.5)
    elif b.startswith('@AUTHORS'):
        p = doc.add_paragraph(); add_rich(p, b[9:].strip()); fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    elif b.startswith('@AFFIL'):
        for line in b[7:].strip().split('\n'):
            p = doc.add_paragraph(); add_rich(p, line.strip()); fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, spacing=1.15)
    elif b.startswith('### '): heading(doc, b[4:], 3)
    elif b.startswith('## '): heading(doc, b[3:], 2)
    elif b.startswith('# '):
        h = b[2:]
        if h == 'Declarations':
            pb = doc.add_paragraph(); pb.add_run().add_break(WD_BREAK.PAGE)
        heading(doc, h, 1)
    elif b.startswith('[[FIG:'):
        figure(doc, int(b[6:-2]))
    elif b.startswith('[[TABLE:'):
        n = int(b[8:-2])
        if n == 1: table1(doc)
        else:
            new_section(doc, landscape=True); line_numbers(doc.sections[-1])
            (table2 if n == 2 else table3)(doc)
            new_section(doc, landscape=False); line_numbers(doc.sections[-1])
    elif b.startswith('•'):
        p = para(doc, b, after=3, hanging=0.6)
    elif b.startswith('**') and ('Keywords' in b[:14]):
        para(doc, b, after=10)
    else:
        para(doc, b, first=0.0 if b.startswith('**') else 1.0)

# references
pb = doc.add_paragraph(); pb.add_run().add_break(WD_BREAK.PAGE)
heading(doc, 'References', 1)
for r in R:
    au, yr, ti, jo, vo, iss, pg, doi, _ = r
    p = doc.add_paragraph()
    add_rich(p, f'{au} ({yr}). {ti}. ')
    jv = f'<i>{jo}</i>'
    if vo: jv += f', <i>{vo}</i>'
    if iss: jv += f'({iss})'
    if pg: jv += f', {pg}'
    add_rich(p, jv + '. ')
    hyperlink(p, f'https://doi.org/{doi}', ('https://doi.org/'+doi).replace('/', '/\u200b').replace('.', '.\u200b').replace('-', '-\u200b'))
    fmt(p, hanging=1.27, after=6, spacing=1.5)

doc.save(OUT)
# word count
body = re.split(r'\n# References', text)[0]
clean = re.sub(r'\[\[.*?\]\]|@\w+|[#*]|\^\{|\}|_\{', ' ', body)
nw = len(clean.split()) + sum(len(re.sub(r'\*|\^\{|\}|_\{', ' ', v[2]).split()) for v in FIGS.values())
print('References:', len(R), '| body words incl. figure captions (excl. tables, refs):', nw)
print('saved', OUT)
