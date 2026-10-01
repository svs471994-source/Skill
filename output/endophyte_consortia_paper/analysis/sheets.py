import random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

F = 'Times New Roman'
YEL = PatternFill('solid', fgColor='FFFF00')
HDR = PatternFill('solid', fgColor='D9E2F3')
EXF = PatternFill('solid', fgColor='EDEDED')
thin = Side(style='thin', color='999999'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def font(b=False, i=False, color='000000', sz=11): return Font(name=F, bold=b, italic=i, color=color, size=sz)
BLUE = '0000FF'

wb = Workbook()

# ---------------- design ----------------
TRT = [(i, s) for i in ('Control', 'IS-06', 'IS-05', 'C4-5') for s in (0, 50, 100, 150)]
REPS = 4
random.seed(2026)
pots = []
for inoc, s in TRT:
    for r in range(1, REPS + 1):
        pots.append((inoc, s, r))
bench = list(range(1, len(pots) + 1)); random.shuffle(bench)
POTS = [dict(id=f'P{i+1:02d}', inoc=p[0], sal=p[1], rep=p[2], trt=f'{p[0]} | {p[1]} mM', pos=bench[i]) for i, p in enumerate(pots)]
EX = dict(id='EX', inoc='EXAMPLE', sal=150, rep=0, trt='EXAMPLE (delete)', pos='')

def title(ws, text, sub=None):
    ws['A1'] = text; ws['A1'].font = font(True, sz=13)
    if sub: ws['A2'] = sub; ws['A2'].font = font(i=True, color='444444')

def table(ws, start_row, groups, rows_fn, example, widths=None):
    """groups: list of (group_label, [(col_header, kind, formula_template_or_None, comment)])
       kind: 'id' (pre-filled design), 'in' (yellow input), 'f' (formula). formula template uses {r} and {c:name}."""
    cols = []
    for g, items in groups:
        for it in items: cols.append((g,) + it)
    # header rows
    gr, hr = start_row, start_row + 1
    c = 1; colidx = {}
    for g, h, kind, fml, com in cols:
        colidx[h] = c; c += 1
    # group header merge
    c = 1
    for g, items in groups:
        n = len(items)
        ws.cell(gr, c, g).font = font(True); ws.cell(gr, c).fill = HDR; ws.cell(gr, c).alignment = Alignment(horizontal='center')
        if n > 1: ws.merge_cells(start_row=gr, start_column=c, end_row=gr, end_column=c + n - 1)
        c += n
    for g, h, kind, fml, com in cols:
        cell = ws.cell(hr, colidx[h], h); cell.font = font(True); cell.fill = HDR
        cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); cell.border = BOX
        if com: cell.comment = Comment(com, 'Data sheet')
    ws.row_dimensions[hr].height = 58
    def ref(name, r): return f'{L(colidx[name])}{r}'
    def write_row(r, p, is_ex):
        vals = rows_fn(p) if not is_ex else example
        for g, h, kind, fml, com in cols:
            cell = ws.cell(r, colidx[h]); cell.border = BOX; cell.font = font()
            if kind == 'id':
                cell.value = {'Pot ID': p['id'], 'Treatment': p['trt'], 'Inoculant': p['inoc'], 'NaCl (mM)': p['sal'], 'Rep': p['rep'], 'Bench position': p['pos']}.get(h, '')
            elif kind == 'in':
                if is_ex and h in vals: cell.value = vals[h]; cell.font = font(color=BLUE)
                if not is_ex: cell.fill = YEL
            elif kind == 'f':
                import re as _re
                stds = sorted(set(_re.findall(r'Std_Curves!\$[A-Z]+\$13', fml)))
                if stds:
                    fml = 'IF(OR(' + ','.join(f'ISNUMBER({x})=FALSE' for x in stds) + '),"",' + fml + ')'
                cell.value = '=' + fml.format(r=r, **{k.replace(' ', '_').replace('(', '').replace(')', '').replace('/', '_').replace('%', 'pct').replace('.', '').replace('-', '_').replace(',', '').replace('+', '').replace('²', '2').replace('µ', 'u').replace('₂', '2').replace('⁺', '').replace('Δ', 'd'): ref(k, r) for k in colidx})
                cell.number_format = '0.000'
            if is_ex: cell.fill = EXF
    r = hr + 1
    write_row(r, EX, True); ws.cell(r, 1).font = font(i=True, color='666666')
    for p in POTS:
        r += 1; write_row(r, p, False)
    ws.freeze_panes = ws.cell(hr + 1, 3)
    for h, ci in colidx.items():
        ws.column_dimensions[L(ci)].width = (widths or {}).get(h, 12 if ci > 2 else 14)
    return colidx, hr + 2, r  # colidx, first data row (after example), last row

def key(name):
    return name.replace(' ', '_').replace('(', '').replace(')', '').replace('/', '_').replace('%', 'pct').replace('.', '').replace('-', '_').replace(',', '').replace('+', '').replace('²', '2').replace('µ', 'u').replace('₂', '2').replace('⁺', '').replace('Δ', 'd')

ID = [('Pot ID', 'id', None, None), ('Treatment', 'id', None, None), ('Inoculant', 'id', None, None), ('NaCl (mM)', 'id', None, None), ('Rep', 'id', None, None)]
def blank_guard(cells, expr):
    cond = 'OR(' + ','.join(f'{{{key(c)}}}=""' for c in cells) + ')'
    return f'IF({cond},"",{expr})'

# ---------------- 0 README ----------------
ws = wb.active; ws.title = 'README'
title(ws, 'Data-recording workbook: Option A (full mechanistic trial) and Option B (minimum add-on)',
      'Rice cv. IR-64 | endophyte treatments Control, IS-06, IS-05, C4-5 | NaCl 0, 50, 100 and 150 mM (same 4-tier format as the first trial) | 4 replicate pots each (64 pots)')
rows = [
 ('HOW TO USE', ''),
 ('Yellow cells', 'Enter your raw readings here (one row per pot; enter the mean of the 3 plants in the pot unless the column says otherwise).'),
 ('White cells with numbers', 'Formulas. Do not type over them; they calculate automatically from your readings.'),
 ('Grey row "EX"', 'Worked example with realistic dummy values to show the expected format. Delete or ignore it; it is excluded from the Summary sheet.'),
 ('Blue text', 'Example values only (not real data).'),
 ('Column headers', 'Hover over a header with a red corner for the method, wavelength and units.'),
 ('', ''),
 ('WHICH SHEETS TO USE', ''),
 ('Option A (full)', 'Layout, A1_Growth_Intervals, A2_Growth_Harvest, A3_Pigments, A4_Osmo_Oxidative, A5_Enzymes, A6_Ions, A7_GasExchange (optional), Std_Curves, Summary.'),
 ('Option B (minimum)', 'Layout, B_Minimal, Std_Curves (proline block), Summary. Same 64 pots; B is a subset of A, so if you run A you do not need B separately.'),
 ('Instruments', 'See sheet "Instruments": what the progress report confirms you have, and what you must confirm or outsource.'),
 ('Timeline', 'See sheet "Timeline" for the day-by-day plan (about 45 days in the polyhouse + 2-3 weeks of lab assays).'),
 ('', ''),
 ('IMPORTANT DESIGN NOTES', ''),
 ('Replicates', 'Record every replicate pot separately. Do NOT enter only means: replicate-level data are what allow ANOVA, MANOVA and correct SDs.'),
 ('Salt timing', 'Write down the exact day saline irrigation starts (sheet Timeline). In the earlier trial the methods said 15 DAS but germination fell with salt, which reviewers will question.'),
 ('Fresh vs dry basis', 'Biochemical values are calculated per g fresh weight (FW) for leaf assays and per g dry weight (DW) for ions, as is standard. Record the exact tissue mass for every sample.'),
 ('Same leaf, same day', 'Take RWC, electrolyte leakage, pigments, MDA, H2O2, proline and enzyme samples from comparable fully expanded leaves on the same day, at the same time of morning.'),
 ('Enzymes', 'Extract on ice in cold phosphate buffer and assay the same day (or store at -20 °C / -80 °C if available). Enter protein (mg/mL) of the SAME enzyme extract to get specific activity.'),
 ('Randomization', 'Bench positions in "Layout" were randomized (completely randomized design, random seed 2026). Place pots accordingly and re-randomize weekly if possible.'),
]
for i, (a, b) in enumerate(rows, start=4):
    ws.cell(i, 1, a).font = font(bool(a) and (b == '' or a.isupper()))
    ws.cell(i, 2, b).font = font(); ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical='top')
ws['A5'].fill = YEL; ws['A7'].fill = EXF
ws.column_dimensions['A'].width = 26; ws.column_dimensions['B'].width = 120

# ---------------- Instruments ----------------
ws = wb.create_sheet('Instruments')
title(ws, 'Instruments: confirmed in your progress report vs. needed for the new parameters',
      'Confirmed = explicitly used in the progress report (Jul-Dec 2025). Confirm = not mentioned in the report; check with your lab / department.')
hdr = ['Parameter', 'Option', 'Method (reference)', 'Instrument needed', 'In progress report?', 'If not available']
data = [
 ('Plant height, leaf number', 'A', 'Ruler / count at 15, 25, 35, 45 DAS', 'Scale, ruler', 'Confirmed (root/shoot lengths measured)', '-'),
 ('Shoot & root fresh weight', 'A, B', 'Gravimetric', 'Analytical balance', 'Confirmed', '-'),
 ('Shoot & root dry weight', 'A, B', 'Oven-dry 70 °C, 72 h to constant weight', 'Hot-air oven', 'Confirm (not mentioned)', 'Use the glassware-drying oven of any microbiology lab; 60-70 °C is enough'),
 ('Leaf area', 'A', 'Scan/photograph leaves on a white sheet with a scale; measure in ImageJ (free)', 'Flatbed scanner or phone camera', 'Confirm', 'Length x max width x 0.75 (rice leaf factor) as a fallback'),
 ('Root volume, surface area, branching', 'A (optional)', 'Scan washed roots in water tray; RhizoVision Explorer (free) or WinRHIZO', 'Flatbed scanner (600 dpi)', 'Confirm', 'Root volume by water displacement in a measuring cylinder; drop surface area/branching and state as limitation'),
 ('Chlorophyll a, b, total, carotenoids', 'A', 'DMSO 60 °C (Hiscox & Israelstam 1979); Arnon (1949) as before; carotenoids Wellburn (1994) DMSO equation at 480 nm', 'UV-Vis spectrophotometer, water bath', 'Confirmed', '-'),
 ('Relative water content', 'A', 'FW, turgid weight (4 h in water, dark), DW (Barrs & Weatherley 1962)', 'Analytical balance, oven', 'Balance confirmed; oven confirm', '-'),
 ('Electrolyte leakage', 'A', 'EC of leaf discs in deionized water, before and after boiling (Lutts et al. 1996)', 'Conductivity (EC) meter', 'Confirm (not mentioned)', 'Borrow from a soil science / chemistry lab; it is a small handheld meter'),
 ('Proline', 'A, B', 'Acid-ninhydrin, toluene phase, A520 (Bates et al. 1973)', 'UV-Vis spectrophotometer, water bath, centrifuge', 'Confirmed', '-'),
 ('Soluble sugars, soluble protein', 'A', 'Anthrone A620; Lowry A660 (as before)', 'UV-Vis spectrophotometer', 'Confirmed', '-'),
 ('MDA (lipid peroxidation)', 'A, B', 'TBA, A532 - A600, ε = 155 mM-1 cm-1 (Heath & Packer 1968)', 'UV-Vis spectrophotometer, water bath, centrifuge', 'Confirmed', '-'),
 ('H2O2', 'A', 'KI method, A390, H2O2 standard curve (Velikova et al. 2000)', 'UV-Vis spectrophotometer, centrifuge', 'Confirmed', '-'),
 ('SOD', 'A', 'NBT photoreduction, A560; 1 U = 50% inhibition (Beauchamp & Fridovich 1971)', 'UV-Vis spectrophotometer, cold centrifuge, light box', 'Spectrophotometer confirmed; cold centrifuge confirm', 'Keep extracts on ice; run within 2 h'),
 ('CAT', 'A', 'H2O2 decomposition, ΔA240/min, ε = 43.6 M-1 cm-1 (Aebi 1984)', 'UV-Vis spectrophotometer with UV range + quartz cuvette', 'Confirm UV range and quartz cuvettes', 'Without quartz cuvettes, the 240 nm and 290 nm assays (CAT, APX) cannot be run'),
 ('APX', 'A', 'Ascorbate oxidation, ΔA290/min, ε = 2.8 mM-1 cm-1 (Nakano & Asada 1981)', 'UV-Vis spectrophotometer (UV) + quartz cuvette', 'Confirm', 'As above'),
 ('POD', 'A', 'Guaiacol, ΔA470/min, ε = 26.6 mM-1 cm-1 (Chance & Maehly 1955)', 'UV-Vis spectrophotometer', 'Confirmed', '-'),
 ('Na+, K+', 'A, B', 'Dry tissue digested (HNO3:HClO4) or dry-ashed; flame photometry', 'Flame photometer', 'Confirm (not mentioned)', 'Outsource to a soil-testing lab or university central instrumentation facility (low cost per sample). Essential for the salinity mechanism.'),
 ('Ca2+, Mg2+', 'A', 'Same digest; EDTA titration (no instrument) or AAS', 'Burette (titration) or AAS', 'Confirm', 'EDTA titration needs only glassware and reagents'),
 ('Photosynthesis, gs, transpiration', 'A (optional)', 'Infrared gas analyser on youngest fully expanded leaf, 9-11 am', 'IRGA (e.g. LI-6800)', 'Not in report', 'Omit and justify as a limitation if no IRGA'),
 ('Chlorophyll fluorescence Fv/Fm', 'A (optional)', '30 min dark adaptation; Fv/Fm', 'Fluorometer (e.g. Handy PEA, MINI-PAM)', 'Not in report', 'Omit and justify as a limitation'),
 ('Incubation, culture, inoculum', 'A, B', 'Nutrient broth, 37 °C; seed bacterization 30 min, ~10^8 CFU/mL', 'Bacteriological incubator, laminar flow, autoclave', 'Confirmed (incubator)', '-'),
]
for j, h in enumerate(hdr, 1):
    c = ws.cell(4, j, h); c.font = font(True); c.fill = HDR; c.border = BOX; c.alignment = Alignment(wrap_text=True)
for i, row in enumerate(data, 5):
    for j, v in enumerate(row, 1):
        c = ws.cell(i, j, v); c.font = font(); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical='top')
        if j == 5 and v.startswith('Confirm'): c.fill = YEL
for j, w in enumerate([30, 11, 52, 34, 26, 48], 1): ws.column_dimensions[L(j)].width = w
ws.cell(len(data) + 6, 1, 'Method references are given for planning; DOIs will be verified when the methods are added to the manuscript.').font = font(i=True)

# ---------------- Layout ----------------
ws = wb.create_sheet('Layout')
title(ws, 'Pot layout: 4 inoculation treatments x 4 NaCl levels (0, 50, 100, 150 mM) x 4 replicates = 64 pots (CRD)',
      'Fill the yellow cells on the day of each event. Bench positions randomized with seed 2026.')
hdr = ['Pot ID', 'Treatment', 'Inoculant', 'NaCl (mM)', 'Rep', 'Bench position', 'Sowing date', 'Seeds sown', 'Seedlings emerged (7 DAS)', 'Emergence %', 'Plants after thinning', 'Plants surviving at harvest', 'Survival %', 'Notes']
for j, h in enumerate(hdr, 1):
    c = ws.cell(4, j, h); c.font = font(True); c.fill = HDR; c.border = BOX; c.alignment = Alignment(wrap_text=True, horizontal='center')
ws.row_dimensions[4].height = 45
for i, p in enumerate(POTS, 5):
    vals = [p['id'], p['trt'], p['inoc'], p['sal'], p['rep'], p['pos']]
    for j, v in enumerate(vals, 1): ws.cell(i, j, v)
    for j in (7, 8, 9, 11, 12, 14): ws.cell(i, j).fill = YEL
    ws.cell(i, 10, f'=IF(OR(H{i}="",I{i}=""),"",I{i}/H{i}*100)').number_format = '0.0'
    ws.cell(i, 13, f'=IF(OR(K{i}="",L{i}=""),"",L{i}/K{i}*100)').number_format = '0.0'
    for j in range(1, 15): ws.cell(i, j).border = BOX; ws.cell(i, j).font = font()
for j, w in enumerate([8, 18, 10, 9, 5, 9, 12, 8, 12, 11, 11, 11, 10, 24], 1): ws.column_dimensions[L(j)].width = w
ws.freeze_panes = 'C5'

# ---------------- A1 intervals ----------------
ws = wb.create_sheet('A1_Growth_Intervals')
title(ws, 'Option A | Growth at defined intervals (mean of the 3 plants per pot)', 'DAS = days after sowing.')
groups = [('Design', ID)]
for das in (15, 25, 35, 45):
    groups.append((f'{das} DAS', [(f'Height {das} DAS (cm)', 'in', None, 'Base of shoot to tip of the tallest leaf, cm'),
                                  (f'Leaves {das} DAS (no.)', 'in', None, 'Number of fully emerged leaves per plant')]))
groups.append(('Growth rate', [('Height gain 15-45 DAS (cm/day)', 'f', blank_guard(['Height 15 DAS (cm)', 'Height 45 DAS (cm)'], '({Height_45_DAS_cm}-{Height_15_DAS_cm})/30'), None)]))
ex = {'Height 15 DAS (cm)': 14.2, 'Leaves 15 DAS (no.)': 3, 'Height 25 DAS (cm)': 19.8, 'Leaves 25 DAS (no.)': 4, 'Height 35 DAS (cm)': 24.1, 'Leaves 35 DAS (no.)': 5, 'Height 45 DAS (cm)': 28.6, 'Leaves 45 DAS (no.)': 6}
table(ws, 4, groups, lambda p: {}, ex)

# ---------------- A2 harvest ----------------
ws = wb.create_sheet('A2_Growth_Harvest')
title(ws, 'Option A | Harvest (45 DAS): biomass, roots, water status', 'Per pot: enter the mean of the 3 plants (biomass per plant).')
groups = [('Design', ID),
 ('Shoot', [('Leaf area (cm2/plant)', 'in', None, 'ImageJ on scanned leaves with a scale bar'),
            ('Shoot FW (g)', 'in', None, None), ('Shoot DW (g)', 'in', None, '70 °C, 72 h, to constant weight')]),
 ('Root', [('Root length (cm)', 'in', None, 'Longest root'), ('Root FW (g)', 'in', None, None), ('Root DW (g)', 'in', None, None),
           ('Root/shoot ratio (DW)', 'f', blank_guard(['Root DW (g)', 'Shoot DW (g)'], '{Root_DW_g}/{Shoot_DW_g}'), 'Root DW / shoot DW'),
           ('Root volume (cm3)', 'in', None, 'Optional: imaging or water displacement'), ('Root surface area (cm2)', 'in', None, 'Optional: imaging'),
           ('Root tips (no.)', 'in', None, 'Optional: imaging'), ('Root forks / branches (no.)', 'in', None, 'Optional: imaging')]),
 ('Relative water content', [('RWC leaf FW (g)', 'in', None, 'Leaf segment fresh weight'), ('RWC turgid wt (g)', 'in', None, 'After 4 h floating on distilled water in dark'),
           ('RWC dry wt (g)', 'in', None, 'After oven drying'),
           ('RWC (%)', 'f', blank_guard(['RWC leaf FW (g)', 'RWC turgid wt (g)', 'RWC dry wt (g)'], '({RWC_leaf_FW_g}-{RWC_dry_wt_g})/({RWC_turgid_wt_g}-{RWC_dry_wt_g})*100'), '(FW-DW)/(TW-DW) x 100')]),
 ('Electrolyte leakage', [('EC1 (µS/cm)', 'in', None, '10 leaf discs in 20 mL deionized water, 2 h, room temp'), ('EC2 (µS/cm)', 'in', None, 'Same tubes after boiling 20 min and cooling'),
           ('Electrolyte leakage (%)', 'f', blank_guard(['EC1 (µS/cm)', 'EC2 (µS/cm)'], '{EC1_uS_cm}/{EC2_uS_cm}*100'), 'EC1/EC2 x 100'),
           ('Membrane stability index (%)', 'f', blank_guard(['EC1 (µS/cm)', 'EC2 (µS/cm)'], '(1-{EC1_uS_cm}/{EC2_uS_cm})*100'), '(1 - EC1/EC2) x 100')]),
]
ex = {'Leaf area (cm2/plant)': 42.5, 'Shoot FW (g)': 3.12, 'Shoot DW (g)': 0.61, 'Root length (cm)': 9.4, 'Root FW (g)': 1.05, 'Root DW (g)': 0.18,
      'Root volume (cm3)': 1.1, 'Root surface area (cm2)': 38.0, 'Root tips (no.)': 410, 'Root forks / branches (no.)': 650,
      'RWC leaf FW (g)': 0.250, 'RWC turgid wt (g)': 0.292, 'RWC dry wt (g)': 0.061, 'EC1 (µS/cm)': 48.0, 'EC2 (µS/cm)': 210.0}
table(ws, 4, groups, lambda p: {}, ex)

# ---------------- Std curves ----------------
sc = wb.create_sheet('Std_Curves')
title(sc, 'Standard curves (absorbance = slope x concentration + intercept)', 'Enter absorbance of each standard. Slope, intercept and R² update automatically and feed the assay sheets.')
blocks = [('Proline (µg/mL) at 520 nm', [0, 5, 10, 20, 30, 40]), ('H2O2 (µmol/mL) at 390 nm', [0, 0.1, 0.2, 0.4, 0.6, 0.8]),
          ('BSA protein (µg/mL) at 660 nm', [0, 40, 80, 120, 160, 200]), ('Glucose (µg/mL) at 620 nm', [0, 20, 40, 60, 80, 100])]
STD = {}
col = 1
for name, concs in blocks:
    sc.cell(4, col, name).font = font(True); sc.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
    sc.cell(5, col, 'Concentration').font = font(True); sc.cell(5, col + 1, 'Absorbance').font = font(True)
    for k, cval in enumerate(concs):
        a = sc.cell(6 + k, col, cval); a.fill = YEL; b = sc.cell(6 + k, col + 1); b.fill = YEL
    rng_x = f'{L(col)}6:{L(col)}{5 + len(concs)}'; rng_y = f'{L(col + 1)}6:{L(col + 1)}{5 + len(concs)}'
    for k, (lab, fn) in enumerate([('Slope', f'=IFERROR(SLOPE({rng_y},{rng_x}),"")'), ('Intercept', f'=IFERROR(INTERCEPT({rng_y},{rng_x}),"")'), ('R²', f'=IFERROR(RSQ({rng_y},{rng_x}),"")')]):
        sc.cell(13 + k, col, lab).font = font(True); c = sc.cell(13 + k, col + 1, fn); c.number_format = '0.0000'
    STD[name.split(' ')[0]] = (f"Std_Curves!${L(col + 1)}$13", f"Std_Curves!${L(col + 1)}$14")
    sc.column_dimensions[L(col)].width = 16; sc.column_dimensions[L(col + 1)].width = 13
    col += 3
sc.cell(17, 1, 'Check: R² should be ≥ 0.99. Re-run any standard curve with each new batch of reagent.').font = font(i=True)
PRO_S, PRO_I = STD['Proline']; H2_S, H2_I = STD['H2O2']; BSA_S, BSA_I = STD['BSA']; GLU_S, GLU_I = STD['Glucose']

# ---------------- A3 pigments ----------------
ws = wb.create_sheet('A3_Pigments')
title(ws, 'Option A | Photosynthetic pigments (DMSO extraction, 60 °C, 30 min)', 'Chl by Arnon (1949) equations as in the earlier trial; carotenoids by the Wellburn (1994) DMSO equation at 480 nm.')
groups = [('Design', ID),
 ('Sample', [('Leaf FW (g)', 'in', None, 'Fresh leaf mass extracted'), ('DMSO volume (mL)', 'in', None, 'Final extract volume')]),
 ('Absorbance', [('A663', 'in', None, None), ('A645', 'in', None, None), ('A480', 'in', None, 'Use 480 nm for carotenoids in DMSO')]),
 ('Concentration in extract (mg/L)', [
    ('Chl a (mg/L)', 'f', blank_guard(['A663', 'A645'], '12.7*{A663}-2.69*{A645}'), None),
    ('Chl b (mg/L)', 'f', blank_guard(['A663', 'A645'], '22.9*{A645}-4.68*{A663}'), None),
    ('Carotenoids (mg/L)', 'f', blank_guard(['A480', 'A663', 'A645'], '(1000*{A480}-1.29*{Chl_a_mg_L}-53.78*{Chl_b_mg_L})/220'), None)]),
 ('Content (mg/g FW)', [
    ('Chl a (mg/g FW)', 'f', blank_guard(['Chl a (mg/L)', 'Leaf FW (g)', 'DMSO volume (mL)'], '{Chl_a_mg_L}*{DMSO_volume_mL}/(1000*{Leaf_FW_g})'), 'C x V / (1000 x W)'),
    ('Chl b (mg/g FW)', 'f', blank_guard(['Chl b (mg/L)', 'Leaf FW (g)', 'DMSO volume (mL)'], '{Chl_b_mg_L}*{DMSO_volume_mL}/(1000*{Leaf_FW_g})'), None),
    ('Total Chl (mg/g FW)', 'f', blank_guard(['Chl a (mg/g FW)', 'Chl b (mg/g FW)'], '{Chl_a_mg_g_FW}+{Chl_b_mg_g_FW}'), 'Chl a + Chl b'),
    ('Carotenoids (mg/g FW)', 'f', blank_guard(['Carotenoids (mg/L)'], '{Carotenoids_mg_L}*{DMSO_volume_mL}/(1000*{Leaf_FW_g})'), None),
    ('Chl a/b ratio', 'f', blank_guard(['Chl a (mg/g FW)', 'Chl b (mg/g FW)'], '{Chl_a_mg_g_FW}/{Chl_b_mg_g_FW}'), None)])]
ex = {'Leaf FW (g)': 0.10, 'DMSO volume (mL)': 5, 'A663': 0.652, 'A645': 0.248, 'A480': 0.410}
table(ws, 4, groups, lambda p: {}, ex)

# ---------------- A4 osmo + oxidative ----------------
ws = wb.create_sheet('A4_Osmo_Oxidative')
title(ws, 'Option A | Osmolytes and oxidative-stress markers (leaf, fresh weight basis)', 'Standard curves are taken from sheet Std_Curves.')
groups = [('Design', ID),
 ('Proline (Bates et al. 1973)', [('Pro leaf FW (g)', 'in', None, 'Tissue homogenized in 10 mL 3% sulfosalicylic acid'), ('Pro A520', 'in', None, 'Toluene phase'),
     ('Toluene (mL)', 'in', None, 'Usually 4 mL'),
     ('Proline (µg/mL)', 'f', blank_guard(['Pro A520'], f'({{Pro_A520}}-{PRO_I})/{PRO_S}'), None),
     ('Proline (µmol/g FW)', 'f', blank_guard(['Proline (µg/mL)', 'Toluene (mL)', 'Pro leaf FW (g)'], '({Proline_ug_mL}*{Toluene_mL}/115.13)/({Pro_leaf_FW_g}/5)'), '[(µg/mL x mL toluene)/115.13] / [g sample/5]')]),
 ('MDA (Heath & Packer 1968)', [('MDA leaf FW (g)', 'in', None, None), ('MDA extract (mL)', 'in', None, 'TCA extract volume'),
     ('A532', 'in', None, None), ('A600', 'in', None, 'Non-specific turbidity'),
     ('MDA (nmol/g FW)', 'f', blank_guard(['A532', 'A600', 'MDA leaf FW (g)', 'MDA extract (mL)'], '(({A532}-{A600})/155)*1000*{MDA_extract_mL}/{MDA_leaf_FW_g}'), 'ε = 155 mM-1 cm-1')]),
 ('H2O2 (Velikova et al. 2000)', [('H2O2 leaf FW (g)', 'in', None, None), ('H2O2 extract (mL)', 'in', None, None), ('H2O2 A390', 'in', None, None),
     ('H2O2 (µmol/g FW)', 'f', blank_guard(['H2O2 A390', 'H2O2 leaf FW (g)', 'H2O2 extract (mL)'], f'(({{H2O2_A390}}-{H2_I})/{H2_S})*{{H2O2_extract_mL}}/{{H2O2_leaf_FW_g}}'), None)]),
 ('Soluble sugars (anthrone)', [('Sug leaf FW (g)', 'in', None, None), ('Sug extract (mL)', 'in', None, None), ('Sug dilution factor', 'in', None, None), ('Sug A620', 'in', None, None),
     ('Soluble sugars (mg/g FW)', 'f', blank_guard(['Sug A620', 'Sug leaf FW (g)', 'Sug extract (mL)', 'Sug dilution factor'], f'(({{Sug_A620}}-{GLU_I})/{GLU_S})*{{Sug_extract_mL}}*{{Sug_dilution_factor}}/({{Sug_leaf_FW_g}}*1000)'), '(C x V x DF)/(W x 1000)')]),
 ('Soluble protein (Lowry)', [('Prot leaf FW (g)', 'in', None, None), ('Prot extract (mL)', 'in', None, None), ('Prot A660', 'in', None, None),
     ('Protein in extract (mg/mL)', 'f', blank_guard(['Prot A660'], f'(({{Prot_A660}}-{BSA_I})/{BSA_S})/1000'), None),
     ('Soluble protein (mg/g FW)', 'f', blank_guard(['Protein in extract (mg/mL)', 'Prot extract (mL)', 'Prot leaf FW (g)'], '{Protein_in_extract_mg_mL}*{Prot_extract_mL}/{Prot_leaf_FW_g}'), None)])]
ex = {'Pro leaf FW (g)': 0.50, 'Pro A520': 0.420, 'Toluene (mL)': 4, 'MDA leaf FW (g)': 0.50, 'MDA extract (mL)': 5, 'A532': 0.385, 'A600': 0.042,
      'H2O2 leaf FW (g)': 0.50, 'H2O2 extract (mL)': 5, 'H2O2 A390': 0.310, 'Sug leaf FW (g)': 0.50, 'Sug extract (mL)': 10, 'Sug dilution factor': 1, 'Sug A620': 0.455,
      'Prot leaf FW (g)': 0.50, 'Prot extract (mL)': 5, 'Prot A660': 0.260}
table(ws, 4, groups, lambda p: {}, ex)

# ---------------- A5 enzymes ----------------
ws = wb.create_sheet('A5_Enzymes')
title(ws, 'Option A | Antioxidant enzymes (cold phosphate-buffer extract; specific activity per mg protein)',
      'Protein of the enzyme extract: enter mg/mL measured on the SAME extract (Lowry/Bradford). Rates: ΔA per minute from the linear part of the trace.')
groups = [('Design', ID),
 ('Extract', [('Enz leaf FW (g)', 'in', None, None), ('Enz extract (mL)', 'in', None, None), ('Extract protein (mg/mL)', 'in', None, 'Protein of this enzyme extract')]),
 ('SOD (NBT, A560)', [('SOD A560 control', 'in', None, 'Reaction without enzyme, illuminated'), ('SOD A560 sample', 'in', None, None), ('SOD enzyme vol (mL)', 'in', None, None),
     ('SOD inhibition (%)', 'f', blank_guard(['SOD A560 control', 'SOD A560 sample'], '({SOD_A560_control}-{SOD_A560_sample})/{SOD_A560_control}*100'), None),
     ('SOD (U/mg protein)', 'f', blank_guard(['SOD inhibition (%)', 'SOD enzyme vol (mL)', 'Extract protein (mg/mL)'], '({SOD_inhibition_pct}/50)/{SOD_enzyme_vol_mL}/{Extract_protein_mg_mL}'), '1 U = 50% inhibition of NBT reduction')]),
 ('CAT (ΔA240)', [('CAT ΔA240/min', 'in', None, 'Decrease in absorbance per minute'), ('CAT reaction vol (mL)', 'in', None, None), ('CAT enzyme vol (mL)', 'in', None, None),
     ('CAT (µmol H2O2/min/mg protein)', 'f', blank_guard(['CAT ΔA240/min', 'CAT reaction vol (mL)', 'CAT enzyme vol (mL)', 'Extract protein (mg/mL)'], '({CAT_dA240_min}/0.0436)*{CAT_reaction_vol_mL}/{CAT_enzyme_vol_mL}/{Extract_protein_mg_mL}'), 'ε = 43.6 M-1 cm-1 = 0.0436 mM-1 cm-1')]),
 ('APX (ΔA290)', [('APX ΔA290/min', 'in', None, None), ('APX reaction vol (mL)', 'in', None, None), ('APX enzyme vol (mL)', 'in', None, None),
     ('APX (µmol ASA/min/mg protein)', 'f', blank_guard(['APX ΔA290/min', 'APX reaction vol (mL)', 'APX enzyme vol (mL)', 'Extract protein (mg/mL)'], '({APX_dA290_min}/2.8)*{APX_reaction_vol_mL}/{APX_enzyme_vol_mL}/{Extract_protein_mg_mL}'), 'ε = 2.8 mM-1 cm-1')]),
 ('POD (guaiacol, ΔA470)', [('POD ΔA470/min', 'in', None, 'Increase in absorbance per minute'), ('POD reaction vol (mL)', 'in', None, None), ('POD enzyme vol (mL)', 'in', None, None),
     ('POD (µmol/min/mg protein)', 'f', blank_guard(['POD ΔA470/min', 'POD reaction vol (mL)', 'POD enzyme vol (mL)', 'Extract protein (mg/mL)'], '({POD_dA470_min}/26.6)*{POD_reaction_vol_mL}/{POD_enzyme_vol_mL}/{Extract_protein_mg_mL}'), 'ε = 26.6 mM-1 cm-1 (tetraguaiacol)')])]
ex = {'Enz leaf FW (g)': 0.50, 'Enz extract (mL)': 5, 'Extract protein (mg/mL)': 1.20, 'SOD A560 control': 0.820, 'SOD A560 sample': 0.410, 'SOD enzyme vol (mL)': 0.05,
      'CAT ΔA240/min': 0.045, 'CAT reaction vol (mL)': 3, 'CAT enzyme vol (mL)': 0.1, 'APX ΔA290/min': 0.030, 'APX reaction vol (mL)': 3, 'APX enzyme vol (mL)': 0.1,
      'POD ΔA470/min': 0.120, 'POD reaction vol (mL)': 3, 'POD enzyme vol (mL)': 0.05}
table(ws, 4, groups, lambda p: {}, ex)

# ---------------- A6 ions ----------------
def ion_groups(prefix_list):
    g = []
    for tis in prefix_list:
        t = tis[:1]
        g.append((f'{tis} digest', [(f'{t} DW digested (g)', 'in', None, 'Oven-dry ground tissue'), (f'{t} digest vol (mL)', 'in', None, 'Final volume after digestion/ashing'),
                                    (f'{t} dilution factor', 'in', None, 'Further dilution before reading; 1 if none')]))
        items = []
        for ion, mw in [('Na', 22.99), ('K', 39.10), ('Ca', 40.08), ('Mg', 24.31)]:
            if prefix_list is B_TIS and ion in ('Ca', 'Mg'): continue
            items.append((f'{t} {ion} reading (mg/L)', 'in', None, 'Instrument reading of the digest (ppm = mg/L)'))
        for ion in (['Na', 'K'] if prefix_list is B_TIS else ['Na', 'K', 'Ca', 'Mg']):
            rd = f'{t} {ion} reading (mg/L)'
            items.append((f'{t} {ion} (mg/g DW)', 'f', blank_guard([rd, f'{t} DW digested (g)', f'{t} digest vol (mL)', f'{t} dilution factor'],
                         f'{{{key(rd)}}}*{{{key(t + " digest vol (mL)")}}}/1000*{{{key(t + " dilution factor")}}}/{{{key(t + " DW digested (g)")}}}'), 'reading x volume/1000 x dilution / DW'))
        items.append((f'{t} K/Na ratio (molar)', 'f', blank_guard([f'{t} Na (mg/g DW)', f'{t} K (mg/g DW)'], f'({{{key(t + " K (mg/g DW)")}}}/39.10)/({{{key(t + " Na (mg/g DW)")}}}/22.99)'), 'Molar ratio: (K/39.10)/(Na/22.99)'))
        g.append((f'{tis} ions', items))
    return g
A_TIS = ['Leaf', 'Root']; B_TIS = ['Leaf ', 'Root ']
ws = wb.create_sheet('A6_Ions')
title(ws, 'Option A | Ion homeostasis in leaves and roots (dry weight basis)', 'Na+, K+ by flame photometry; Ca2+, Mg2+ by EDTA titration or AAS (enter the result as mg/L of the digest).')
groups = [('Design', ID)] + ion_groups(A_TIS)
ex = {'L DW digested (g)': 0.20, 'L digest vol (mL)': 50, 'L dilution factor': 1, 'L Na reading (mg/L)': 48.0, 'L K reading (mg/L)': 96.0, 'L Ca reading (mg/L)': 22.0, 'L Mg reading (mg/L)': 11.0,
      'R DW digested (g)': 0.20, 'R digest vol (mL)': 50, 'R dilution factor': 1, 'R Na reading (mg/L)': 70.0, 'R K reading (mg/L)': 52.0, 'R Ca reading (mg/L)': 15.0, 'R Mg reading (mg/L)': 8.0}
table(ws, 4, groups, lambda p: {}, ex)

# ---------------- A7 gas exchange ----------------
ws = wb.create_sheet('A7_GasExchange')
title(ws, 'Option A (optional) | Gas exchange and chlorophyll fluorescence', 'Only if an IRGA / fluorometer is available. Measure the youngest fully expanded leaf, 09:00-11:00 h.')
groups = [('Design', ID),
 ('Gas exchange (IRGA)', [('Pn (µmol CO2/m2/s)', 'in', None, 'Net photosynthetic rate'), ('gs (mol H2O/m2/s)', 'in', None, 'Stomatal conductance'),
                          ('E (mmol H2O/m2/s)', 'in', None, 'Transpiration rate'), ('Ci (µmol/mol)', 'in', None, 'Intercellular CO2'),
                          ('WUE (Pn/E)', 'f', blank_guard(['Pn (µmol CO2/m2/s)', 'E (mmol H2O/m2/s)'], '{Pn_umol_CO2_m2_s}/{E_mmol_H2O_m2_s}'), 'Instantaneous water-use efficiency')]),
 ('Fluorescence (30 min dark)', [('F0', 'in', None, None), ('Fm', 'in', None, None),
                                 ('Fv/Fm', 'f', blank_guard(['F0', 'Fm'], '({Fm}-{F0})/{Fm}'), 'Maximum PSII efficiency')])]
ex = {'Pn (µmol CO2/m2/s)': 14.8, 'gs (mol H2O/m2/s)': 0.32, 'E (mmol H2O/m2/s)': 4.6, 'Ci (µmol/mol)': 265, 'F0': 410, 'Fm': 2050}
table(ws, 4, groups, lambda p: {}, ex)

# ---------------- B minimal ----------------
ws = wb.create_sheet('B_Minimal')
title(ws, 'Option B | Minimum add-on: shoot/root dry weight, MDA, proline, leaf and root Na+/K+',
      'Same 64 pots and harvest day as Option A. Proline standard curve from Std_Curves.')
groups = [('Design', ID),
 ('Biomass', [('Shoot DW (g)', 'in', None, '70 °C, 72 h'), ('Root DW (g)', 'in', None, None),
              ('Root/shoot ratio', 'f', blank_guard(['Shoot DW (g)', 'Root DW (g)'], '{Root_DW_g}/{Shoot_DW_g}'), None)]),
 ('MDA', [('MDA leaf FW (g)', 'in', None, None), ('MDA extract (mL)', 'in', None, None), ('A532', 'in', None, None), ('A600', 'in', None, None),
          ('MDA (nmol/g FW)', 'f', blank_guard(['A532', 'A600', 'MDA leaf FW (g)', 'MDA extract (mL)'], '(({A532}-{A600})/155)*1000*{MDA_extract_mL}/{MDA_leaf_FW_g}'), 'ε = 155 mM-1 cm-1')]),
 ('Proline', [('Pro leaf FW (g)', 'in', None, None), ('Pro A520', 'in', None, None), ('Toluene (mL)', 'in', None, None),
              ('Proline (µg/mL)', 'f', blank_guard(['Pro A520'], f'({{Pro_A520}}-{PRO_I})/{PRO_S}'), None),
              ('Proline (µmol/g FW)', 'f', blank_guard(['Proline (µg/mL)', 'Toluene (mL)', 'Pro leaf FW (g)'], '({Proline_ug_mL}*{Toluene_mL}/115.13)/({Pro_leaf_FW_g}/5)'), None)])] + ion_groups(B_TIS)
ex = {'Shoot DW (g)': 0.61, 'Root DW (g)': 0.18, 'MDA leaf FW (g)': 0.50, 'MDA extract (mL)': 5, 'A532': 0.385, 'A600': 0.042, 'Pro leaf FW (g)': 0.50, 'Pro A520': 0.420, 'Toluene (mL)': 4,
      'L DW digested (g)': 0.20, 'L digest vol (mL)': 50, 'L dilution factor': 1, 'L Na reading (mg/L)': 48.0, 'L K reading (mg/L)': 96.0,
      'R DW digested (g)': 0.20, 'R digest vol (mL)': 50, 'R dilution factor': 1, 'R Na reading (mg/L)': 70.0, 'R K reading (mg/L)': 52.0}
bcols, _, _ = table(ws, 4, groups, lambda p: {}, ex)

# ---------------- Summary ----------------
sm = wb.create_sheet('Summary')
title(sm, 'Summary: treatment means and n (updates automatically; EX rows excluded)', 'Use these means for a quick look only; the full statistics (ANOVA/MANOVA) will be run on the replicate rows.')
traits = [('B_Minimal', 'Shoot DW (g)'), ('B_Minimal', 'Root DW (g)'), ('B_Minimal', 'Root/shoot ratio'), ('B_Minimal', 'MDA (nmol/g FW)'), ('B_Minimal', 'Proline (µmol/g FW)'),
          ('B_Minimal', 'L Na (mg/g DW)'), ('B_Minimal', 'L K (mg/g DW)'), ('B_Minimal', 'L K/Na ratio (molar)'), ('B_Minimal', 'R Na (mg/g DW)'), ('B_Minimal', 'R K/Na ratio (molar)'),
          ('A2_Growth_Harvest', 'RWC (%)'), ('A2_Growth_Harvest', 'Electrolyte leakage (%)'), ('A3_Pigments', 'Total Chl (mg/g FW)'), ('A4_Osmo_Oxidative', 'H2O2 (µmol/g FW)'),
          ('A5_Enzymes', 'SOD (U/mg protein)'), ('A5_Enzymes', 'CAT (µmol H2O2/min/mg protein)'), ('A5_Enzymes', 'APX (µmol ASA/min/mg protein)'), ('A5_Enzymes', 'POD (µmol/min/mg protein)'),
          ('A6_Ions', 'L K/Na ratio (molar)'), ('A6_Ions', 'R K/Na ratio (molar)')]
# find columns
colmap = {}
for s in set(t[0] for t in traits):
    w = wb[s]
    for c in range(1, w.max_column + 1):
        v = w.cell(5, c).value
        if v: colmap[(s, v)] = L(c)
sm.cell(4, 1, 'Treatment').font = font(True); sm.cell(4, 1).fill = HDR
sm.cell(4, 2, 'n (B_Minimal rows with MDA)').font = font(True); sm.cell(4, 2).fill = HDR
for j, (s, t) in enumerate(traits, 3):
    c = sm.cell(4, j, f'{t} [{s}]'); c.font = font(True); c.fill = HDR; c.alignment = Alignment(wrap_text=True)
    sm.column_dimensions[L(j)].width = 14
sm.row_dimensions[4].height = 75
last = 6 + len(POTS)
for i, (inoc, s) in enumerate(TRT, 5):
    lab = f'{inoc} | {s} mM'; sm.cell(i, 1, lab).font = font()
    mcol = colmap[('B_Minimal', 'MDA (nmol/g FW)')]
    sm.cell(i, 2, f'=SUMPRODUCT((B_Minimal!$B$7:$B${last}=$A{i})*(B_Minimal!${mcol}$7:${mcol}${last}<>""))')
    for j, (sh, t) in enumerate(traits, 3):
        col = colmap[(sh, t)]
        c = sm.cell(i, j, f'=IFERROR(AVERAGEIFS({sh}!${col}$7:${col}${last},{sh}!$B$7:$B${last},$A{i}),"")'); c.number_format = '0.000'
sm.column_dimensions['A'].width = 18; sm.column_dimensions['B'].width = 12
sm.freeze_panes = 'B5'

# ---------------- Timeline ----------------
tl = wb.create_sheet('Timeline')
title(tl, 'Timeline (both options; Option B skips the rows marked A only)', 'Enter actual dates in the yellow column.')
steps = [
 ('-7 to -1', 'A, B', 'Revive IS-05, IS-06 (and IS-04 for C4-5) on nutrient agar; purity check; prepare potting mix (soil:cocopeat 1:2), sterilize, fill 64 pots of 2.5 kg'),
 ('0', 'A, B', 'Surface-sterilize IR-64 seeds; bacterize 30 min (~10^8 CFU/mL; check by plate count); sow 10 seeds per pot; record in Layout'),
 ('7', 'A, B', 'Count emergence (Layout)'),
 ('10', 'A, B', 'Thin to 3 uniform plants per pot'),
 ('15', 'A, B', 'START NaCl irrigation (record exact date). 50, 100 and 150 mM NaCl (2.922, 5.844 and 8.766 g/L): 500 mL on alternate days for 30 days; 0 mM pots get the same volume of water. Heights and leaf counts (A1)'),
 ('25, 35', 'A only', 'Heights and leaf counts (A1)'),
 ('40-44', 'A only', 'Gas exchange and Fv/Fm, 09:00-11:00 h (A7), if instruments available'),
 ('45 (harvest, morning)', 'A, B', 'Heights and leaf counts (A1); sample leaves for RWC, electrolyte leakage, pigments, MDA, H2O2, proline, sugars, protein, enzymes (keep on ice); uproot and wash roots'),
 ('45', 'A only', 'Root scanning before drying (A2), leaf area scans'),
 ('45', 'A, B', 'Shoot and root fresh weights; oven-dry at 70 °C for 72 h'),
 ('45-46', 'A, B', 'MDA and proline (same day if possible); A: pigments, H2O2, enzymes on fresh extract'),
 ('48-60', 'A, B', 'Dry weights; grind dry tissue; digest for Na, K (and Ca, Mg for A); flame photometer or outsourced analysis'),
 ('60-65', 'A, B', 'Enter all data; send the workbook back for statistics and manuscript update'),
]
for j, h in enumerate(['Day (DAS)', 'Option', 'Activity', 'Actual date'], 1):
    c = tl.cell(4, j, h); c.font = font(True); c.fill = HDR; c.border = BOX
for i, st in enumerate(steps, 5):
    for j, v in enumerate(st, 1):
        c = tl.cell(i, j, v); c.font = font(); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical='top')
    tl.cell(i, 4).fill = YEL; tl.cell(i, 4).border = BOX
for j, w in enumerate([18, 9, 110, 14], 1): tl.column_dimensions[L(j)].width = w

# fonts for any unset cells
for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            if c.value is not None and (c.font is None or c.font.name != F):
                c.font = Font(name=F, bold=c.font.bold if c.font else False, italic=c.font.italic if c.font else False, color=c.font.color if c.font else None, size=c.font.size if c.font and c.font.size else 11)

wb.save('Data_Recording_Sheets_OptionA_OptionB.xlsx')
print('saved')
