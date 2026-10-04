import re
from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
src=open("Obj4_Perception_manuscript.md").read().split("\n")
doc=Document()
st=doc.styles["Normal"]; st.font.name="Times New Roman"; st.font.size=Pt(11)
for s in doc.sections: s.left_margin=s.right_margin=Cm(2.3); s.top_margin=s.bottom_margin=Cm(2.3)
def runs(par,text):
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)",text):
        if not tok: continue
        if tok.startswith("**"): r=par.add_run(tok[2:-2]); r.bold=True
        elif tok.startswith("*") and len(tok)>2: r=par.add_run(tok[1:-1]); r.italic=True
        else: par.add_run(tok)
i=0
while i<len(src):
    l=src[i]
    if l.startswith("# "): p=doc.add_heading(l[2:],0)
    elif l.startswith("## "): doc.add_heading(l[3:],1)
    elif l.startswith("### "): doc.add_heading(l[4:],2)
    elif l.strip()=="---": pass
    elif l.startswith("!["):
        path=re.search(r"\((.*?)\)",l).group(1)
        doc.add_picture(path,width=Inches(5.6)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif l.startswith("|"):
        rows=[]
        while i<len(src) and src[i].startswith("|"):
            if not re.match(r"^\|[-| ]+\|$",src[i]): rows.append([c.strip() for c in src[i].strip("|").split("|")])
            i+=1
        t=doc.add_table(rows=len(rows),cols=len(rows[0])); t.style="Light Grid Accent 1"
        for a,r in enumerate(rows):
            for b,c in enumerate(r):
                cell=t.cell(a,b); cell.text=""; runs(cell.paragraphs[0],c)
                for rr in cell.paragraphs[0].runs: rr.font.size=Pt(9); rr.bold = rr.bold or a==0
        doc.add_paragraph(); continue
    elif re.match(r"^- ",l): runs(doc.add_paragraph(style="List Bullet"),l[2:])
    elif re.match(r"^\d+\. ",l): runs(doc.add_paragraph(style="List Number"),re.sub(r"^\d+\. ","",l))
    elif l.strip(): 
        p=doc.add_paragraph(); runs(p,l)
        if l.startswith("**Fig.") or l.startswith("**Table"): 
            for r in p.runs: r.font.size=Pt(9.5)
    i+=1
doc.save("Obj4_Perception_manuscript.docx")
